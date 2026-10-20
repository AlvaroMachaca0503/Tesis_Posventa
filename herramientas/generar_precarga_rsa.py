# -*- coding: utf-8 -*-
"""Prepara la precarga de RSA (grupos G1-G5 y servicios con tarifas) a partir de
los dos Excel de trabajo y genera:
  - prototipo/capas/r77/ag_rsa_data.js    (datos que trae el botón «Cargar datos de RSA»)
  - prototipo/carga_masiva/*.xlsx          (precarga y plantillas para la carga masiva)

Usa SOLO datos de negocio (reglas de grupo, servicios y tarifas). NO copia
clientes, placas, ODS, usuarios, correos ni el caché de la API: esos son datos
personales y se quedan fuera de la carpeta sincronizada.

Uso: python herramientas/generar_precarga_rsa.py "<Tabla_maestra...xlsx>" "<Base de datos RSA...xlsx>"
"""
import io, json, os, re, sys, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
G15 = ('G1', 'G2', 'G3', 'G4', 'G5')


def hojas(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    return wb


def tabla(ws):
    filas = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    h = [str(c).strip() if c is not None else '' for c in filas[0]]
    return h, filas[1:]


def num(v):
    if v is None or v == '':
        return None
    try:
        f = float(str(v).replace(',', '.'))
        return int(f) if f == int(f) else round(f, 2)
    except ValueError:
        return None


def txt(v):
    return '' if v is None else str(v).strip()


def norm(s):
    return re.sub(r'[^A-Z0-9]', '', txt(s).upper())


def main(maestra, base):
    wm = hojas(maestra)
    h, filas = tabla(wm['Reglas G1-G5'])
    c = {k: i for i, k in enumerate(h)}
    grupos = collections.OrderedDict()
    reglas = []
    for r in filas:
        g = txt(r[c['GRUPO']])
        if g not in G15:
            continue
        grupos.setdefault(g, {'id': g, 'nombre': txt(r[c['NOMBRE_GRUPO']]), 'descripcion': ''})
        reglas.append({
            'grupo': g, 'marca': txt(r[c['MARCA']]), 'modelo': txt(r[c['MODELO']]),
            'anio_desde': num(r[c['ANIO_DESDE']]), 'anio_hasta': num(r[c['ANIO_HASTA']]),
            'codigos_motor': txt(r[c['CODIGOS_MOTOR_VERIFICADOS']]).replace(' ', ''),
            'estado_codigo': txt(r[c['ESTADO_CODIGO']]),
            'aplicacion_automatica': 'SI' if txt(r[c['DECISION_OPERATIVA']]) == 'ASIGNAR GRUPO' else 'NO',
            'confianza': num(r[c['CONFIANZA_REGLA']]),
            'condicion_campo': txt(r[c['CAMPO_ADICIONAL']]), 'condicion_valor': txt(r[c['VALOR_ADICIONAL']]),
            'decision': txt(r[c['DECISION_OPERATIVA']]),
        })
    wb = hojas(base)
    # descripciones de grupo (hoja Grupos de la base)
    hg, fg = tabla(wb['Grupos'])
    cg = {k: i for i, k in enumerate(hg)}
    for r in fg:
        g = txt(r[cg['GRUPO_ID']])
        if g in grupos and not grupos[g]['descripcion']:
            grupos[g]['descripcion'] = txt(r[cg['DESCRIPCION_GRUPO']])
    # servicios y tarifas
    hs, fs = tabla(wb['Servicios'])
    cs = {k: i for i, k in enumerate(hs)}
    servicios = collections.OrderedDict()
    tarifas = {}
    excepciones = []
    idx = collections.defaultdict(list)  # (marca, modelo) -> [(g, desde, hasta)]
    for r in reglas:
        idx[(norm(r['marca']), norm(r['modelo']))].append((r['grupo'], r['anio_desde'], r['anio_hasta']))
    for r in fs:
        cod = txt(r[cs['CODIGO_SERVICIO']])
        if txt(r[cs['ESTADO']]) != 'ACTIVO':
            continue
        sv = servicios.setdefault(cod, {'codigo': cod, 'nombre': txt(r[cs['NOMBRE_SERVICIO']]), 'categoria': txt(r[cs['CATEGORIA']]), 'descripcion': txt(r[cs['REGLA_APLICACION']]), 'requiere_confirmacion': txt(r[cs['REQUIERE_CONFIRMACION']])})
        if not sv['descripcion']:
            sv['descripcion'] = txt(r[cs['REGLA_APLICACION']])
        hh = num(r[cs['HH_MAXIMA']])
        precio = num(r[cs['PRECIO']])
        if txt(r[cs['TIPO_REGLA']]) == 'GRUPO':
            g = txt(r[cs['GRUPO_ID']])
            if g in G15:
                tarifas[(cod, g)] = {'codigo': cod, 'grupo': g, 'hh': hh, 'precio': precio, 'origen': txt(r[cs['ORIGEN_PRECIO']]), 'id_tarifa': txt(r[cs['ID_TARIFA']]), 'codigo_barra': txt(r[cs['CODIGO_BARRA']])}
        else:
            marca, modelo = txt(r[cs['MARCA']]), txt(r[cs['MODELO']])
            d, a = num(r[cs['ANIO_DESDE']]), num(r[cs['ANIO_HASTA']])
            gs = {g for (g, gd, ga) in idx.get((norm(marca), norm(modelo)), []) if (d is None or ga is None or d <= ga) and (a is None or gd is None or a >= gd)}
            for g in sorted(gs):
                excepciones.append({'codigo': cod, 'grupo': g, 'marca': marca, 'modelo': modelo, 'anio_desde': d, 'anio_hasta': a,
                                    'condicion_campo': txt(r[cs['CAMPO_CONDICION_1']]), 'condicion_valor': txt(r[cs['VALOR_CONDICION_1']]),
                                    'hh': hh, 'precio': precio, 'origen': txt(r[cs['ORIGEN_PRECIO']])})
    datos = {
        'fuente': 'Tabla maestra G1-G5 y Base de datos RSA (estandarización de servicios): solo reglas, servicios y tarifas',
        'grupos': list(grupos.values()), 'reglas': reglas,
        'servicios': list(servicios.values()), 'tarifas': list(tarifas.values()), 'excepciones': excepciones,
    }
    ruta_js = os.path.join(RAIZ, 'prototipo', 'capas', 'r77', 'ag_rsa_data.js')
    with io.open(ruta_js, 'w', encoding='utf-8', newline='\n') as f:
        f.write('/* Generado por herramientas/generar_precarga_rsa.py: solo reglas de grupo, servicios y tarifas (sin datos personales). */\nconst R77_RSA=' + json.dumps(datos, ensure_ascii=False, separators=(',', ':')) + ';\n')
    # --- Excel
    carpeta = os.path.join(RAIZ, 'prototipo', 'carga_masiva')
    os.makedirs(carpeta, exist_ok=True)
    azul = PatternFill('solid', fgColor='1F3A5F')

    def escribir(ws, cols, filas, anchos=None):
        ws.append(cols)
        for c_ in ws[1]:
            c_.font = Font(bold=True, color='FFFFFF')
            c_.fill = azul
            c_.alignment = Alignment(vertical='center')
        for f_ in filas:
            ws.append(f_)
        ws.freeze_panes = 'A2'
        for i, col in enumerate(cols):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = (anchos or {}).get(col, max(12, min(40, len(col) + 4)))

    def instrucciones(wb, titulo, lineas):
        ws = wb.create_sheet('Instrucciones', 0)
        ws['A1'] = titulo
        ws['A1'].font = Font(bold=True, size=14)
        for i, l in enumerate(lineas, start=3):
            ws.cell(row=i, column=1, value=l)
        ws.column_dimensions['A'].width = 120

    modos = ['Modos de carga (se elige en Configuración › Carga masiva):',
             '  • Solo agregar nuevos: no toca lo que ya existe.',
             '  • Agregar y actualizar: actualiza los campos de lo que coincide y agrega lo nuevo.',
             '  • Agregar y completar vacíos: agrega lo nuevo y solo rellena campos vacíos de lo existente (no pisa lo que ya tiene valor).',
             'Antes de aplicar, la pantalla muestra una vista previa (nuevos, actualizados, sin cambios y errores). No se borra nada.',
             'No cambies los nombres de las columnas ni de las hojas. Una fila por registro.']
    wb1 = openpyxl.Workbook()
    wb1.remove(wb1.active)
    instrucciones(wb1, 'Precarga RSA · grupos G1 a G5 y servicios con tarifas', [
        'Contenido: 5 grupos con sus reglas de asignación (marca + modelo + rango de año), 14 servicios estandarizados,',
        'tarifas (horas hombre y precio) por grupo y excepciones por marca y modelo. Todo es de negocio, sin datos personales.',
        'Tipo de carga en pantalla: «Grupos y reglas» y «Servicios y tarifas» (cada una lee sus hojas).',
        'Reglas con aplicacion_automatica = NO quedan como SUGERENCIA (el sistema propone el grupo pero no lo asigna solo).',
        'El código de motor solo informa; la asignación usa marca + modelo + rango de año.', ''] + modos)
    escribir(wb1.create_sheet('Grupos'), ['grupo', 'nombre', 'descripcion'], [[g['id'], g['nombre'], g['descripcion']] for g in grupos.values()])
    escribir(wb1.create_sheet('Reglas'), ['grupo', 'marca', 'modelo', 'anio_desde', 'anio_hasta', 'codigos_motor', 'aplicacion_automatica', 'confianza', 'condicion_campo', 'condicion_valor', 'estado_codigo'],
             [[r['grupo'], r['marca'], r['modelo'], r['anio_desde'], r['anio_hasta'], r['codigos_motor'], r['aplicacion_automatica'], r['confianza'], r['condicion_campo'], r['condicion_valor'], r['estado_codigo']] for r in reglas])
    escribir(wb1.create_sheet('Servicios'), ['codigo', 'nombre', 'categoria', 'descripcion', 'requiere_confirmacion'],
             [[s['codigo'], s['nombre'], s['categoria'], s['descripcion'], s['requiere_confirmacion']] for s in servicios.values()], {'nombre': 52, 'descripcion': 70})
    escribir(wb1.create_sheet('Tarifas'), ['codigo', 'grupo', 'hh', 'precio', 'origen'], [[t['codigo'], t['grupo'], t['hh'], t['precio'], t['origen']] for t in tarifas.values()], {'origen': 60})
    escribir(wb1.create_sheet('Excepciones'), ['codigo', 'grupo', 'marca', 'modelo', 'anio_desde', 'anio_hasta', 'condicion_campo', 'condicion_valor', 'hh', 'precio', 'origen'],
             [[e['codigo'], e['grupo'], e['marca'], e['modelo'], e['anio_desde'], e['anio_hasta'], e['condicion_campo'], e['condicion_valor'], e['hh'], e['precio'], e['origen']] for e in excepciones], {'origen': 60})
    wb1.save(os.path.join(carpeta, '01_RSA_precarga_grupos_y_servicios_G1_G5.xlsx'))

    plantillas = [
        ('02_plantilla_clientes.xlsx', 'Clientes', 'clientes', ['tipo_documento', 'documento', 'nombre', 'celular', 'correo'], [['DNI', '00000000', 'NOMBRE DE EJEMPLO', '999000000', 'ejemplo@correo.com']], 'La llave es el documento. Usa datos sintéticos en las pruebas.'),
        ('03_plantilla_vehiculos.xlsx', 'Vehiculos', 'vehiculos', ['placa', 'marca', 'modelo', 'anio', 'color', 'vin', 'serie', 'motor', 'aceite', 'documento_cliente'], [['AAA-000', 'KIA', 'PICANTO', 2019, 'GRIS', '', '', 'G4LA', 'Sintético', '00000000']], 'La llave es la placa. aceite: Mineral, Semisintético, Sintético o Full sintético. El grupo y el plan se asignan solos con las reglas.'),
        ('04_plantilla_productos_y_servicios.xlsx', 'Catalogo', 'catalogo', ['codigo', 'nombre', 'tipo', 'categoria', 'precio', 'costo', 'horas_hombre', 'tipo_producto'], [['P0000', 'PRODUCTO DE EJEMPLO', 'producto', 'INSUMOS', 45, 30, '', 'interno']], 'La llave es el código. tipo: producto o servicio. tipo_producto: interno (entra al stock) o externo.'),
        ('05_plantilla_proveedores.xlsx', 'Proveedores', 'proveedores', ['tipo_documento', 'documento', 'nombre', 'celular'], [['RUC', '20000000000', 'PROVEEDOR DE EJEMPLO S.A.C.', '999000000']], 'La llave es el documento.'),
        ('06_plantilla_stock_inicial.xlsx', 'Stock', 'stock', ['codigo', 'local', 'cantidad', 'precio_compra', 'precio_venta'], [['P0000', 'Almacén General', 10, 30, 45]], 'Suma al kárdex como «Inventario inicial» o «Ajuste por carga». local: el nombre del almacén. precio_compra y precio_venta son opcionales (precio por local).'),
        ('07_plantilla_solapamientos.xlsx', 'Solapamientos', 'solapes', ['servicio_base', 'servicio_que_se_reduce', 'reduccion_hh', 'fuente'], [['REF', 'RAD', 0.5, 'Autodata']], 'Códigos del catálogo. La llave es la pareja servicio_base + servicio_que_se_reduce.'),
        ('09_plantilla_ordenes_de_servicio.xlsx', 'Ordenes', 'ordenes', ['codigo_ods', 'fecha', 'hora', 'placa', 'kilometraje', 'estado', 'tipo_atencion', 'motivo', 'sede', 'tecnico_orden', 'asesor', 'tipo_item', 'codigo_item', 'item', 'cantidad', 'precio_unitario', 'costo', 'descuento', 'tecnico_servicio', 'documento_cliente', 'nombre_cliente', 'marca', 'modelo', 'anio'], [['100001', '2026-10-01', '09:30', 'AAA-000', 50000, 'En proceso', 'Mantenimiento', 'Cambio de aceite y filtros', 'union', '', '', 'servicio', 'ACE', 'CAMBIO DE ACEITE DE MOTOR', 0.75, 30, '', 0, '', '00000000', 'NOMBRE DE EJEMPLO', 'KIA', 'PICANTO', 2019]], 'Una fila por ítem; las filas con el mismo codigo_ods forman una orden. fecha: AAAA-MM-DD o DD/MM/AAAA. estado: Recepción, En proceso, Entrega o Cerrada. sede: union o prohogar. tipo_item: servicio o producto. Si la placa no existe se crea el vehículo (y el cliente con documento_cliente y nombre_cliente). Las órdenes abiertas crean su tarea y activan las automatizaciones.'),
        ('08_plantilla_paquetes.xlsx', 'Paquetes', 'paquetes', ['paquete', 'promocion', 'codigo', 'cantidad'], [['Mantenimiento preventivo menor', '10 % por el paquete completo', 'ACE', 1], ['Mantenimiento preventivo menor', '10 % por el paquete completo', 'BUJ', 1]], 'Una fila por ítem. La llave es el nombre del paquete.'),
    ]
    for nombre, hoja, ent, cols, ej, nota in plantillas:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        instrucciones(wb, 'Plantilla de carga masiva · ' + hoja, ['Tipo de carga en pantalla: ' + ent, nota, 'Borra la fila de ejemplo y escribe tus filas. Usa solo datos tuyos o sintéticos.', ''] + modos)
        escribir(wb.create_sheet(hoja), cols, ej)
        wb.save(os.path.join(carpeta, nombre))
    print('grupos', len(grupos), 'reglas', len(reglas), '(auto SI', sum(1 for r in reglas if r['aplicacion_automatica'] == 'SI'), '/ NO', sum(1 for r in reglas if r['aplicacion_automatica'] == 'NO'), ')')
    print('servicios', len(servicios), 'tarifas', len(tarifas), 'excepciones', len(excepciones))
    print('archivos', sorted(os.listdir(carpeta)))
    print('js bytes', os.path.getsize(ruta_js))


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
